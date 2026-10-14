#!/usr/bin/env python3
"""Publish a source-backed Periplus Markdown manuscript in legacy Seven Ninths layout.

A4, Seven Ninths branded header, PAGE / NUMPAGES footer, cover, three-level
clickable paginated TOC. Two-pass pagination uses LibreOffice and PDF outline
page destinations to create stable TOC page numbers even in headless export.
This script formats existing researched text; it does NOT perform web research.
"""
from __future__ import annotations
import argparse
from datetime import date
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from trip_dates import display_trip_dates, publication_stem

BRAND = "Seven Ninths Travel Brochure"
WESTERN_FONT = "Oswald"
EAST_ASIAN_FONT = "Noto Sans CJK SC"
ROOT_NAME = "Seven_Ninths_Travel_Brochure"
MAIN_SECTIONS = {
    "Arrival & Departure", "Accommodation", "Brief Sketch", "Attractions & Schedule",
    "Sightseeing & Adventure", "Delicacy & Shopping", "Recommends & Tips",
}


def add_field(par, instruction: str, fallback='1'):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = par.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'begin'); r._r.append(e)
    r = par.add_run(); e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = instruction; r._r.append(e)
    r = par.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'separate'); r._r.append(e)
    par.add_run(fallback)
    r = par.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)


def toc_field_open(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind in ('begin', 'instruction', 'separate'):
        r = par.add_run()
        if kind == 'instruction':
            x = OxmlElement('w:instrText')
            x.set(qn('xml:space'), 'preserve')
            x.text = ' TOC \\o "1-3" \\h \\z \\u '
        else:
            x = OxmlElement('w:fldChar')
            x.set(qn('w:fldCharType'), kind)
        r._r.append(x)


def toc_field_close(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = par.add_run(); x = OxmlElement('w:fldChar'); x.set(qn('w:fldCharType'), 'end'); r._r.append(x)


def bookmark(par, idx):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(idx+1)); start.set(qn('w:name'), f'Periplus_H_{idx+1}')
    stop = OxmlElement('w:bookmarkEnd'); stop.set(qn('w:id'), str(idx+1))
    par._p.insert(0, start); par._p.append(stop)


def internal_link(par, text, idx):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    x = OxmlElement('w:hyperlink'); x.set(qn('w:anchor'), f'Periplus_H_{idx+1}')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = text
    r.append(t); x.append(r); par._p.append(x)


def external_link(par, txt, url):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    rel = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    x = OxmlElement('w:hyperlink'); x.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '245D88'); pr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    r.append(pr); t = OxmlElement('w:t'); t.text = txt; r.append(t); x.append(r); par._p.append(x)


def style_font(style, western=WESTERN_FONT, east=EAST_ASIAN_FONT, size=10, color=None):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    style.font.name = western; style.font.size = Pt(size)
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn('w:eastAsia'), east)
    if color is not None: style.font.color.rgb = color


def parse_headings(tokens):
    out = []
    first = next(((i, t) for i, t in enumerate(tokens[:-1]) if t.type == 'heading_open'), None)
    cover_index = None
    if first:
        i, t = first
        label = tokens[i+1].content.strip()
        if int(t.tag[1]) == 1 and label not in MAIN_SECTIONS and not label.startswith('Arrival & Departure'):
            cover_index = i
    for i, tok in enumerate(tokens[:-1]):
        if tok.type != 'heading_open' or i == cover_index: continue
        label = tokens[i+1].content.strip()
        if not label or label.lower() in ('content', 'contents', 'table of contents', '\u76ee\u5f55'): continue
        out.append((i, min(int(tok.tag[1]), 4), label))
    return out, cover_index


def inline(par, token, manuscript):
    from docx.shared import Cm
    children = token.children or []
    bold = italic = False
    link = None
    for child in children:
        k = child.type
        if k == 'strong_open': bold = True
        elif k == 'strong_close': bold = False
        elif k == 'em_open': italic = True
        elif k == 'em_close': italic = False
        elif k == 'link_open': link = child.attrGet('href')
        elif k == 'link_close': link = None
        elif k in ('text', 'code_inline', 'html_inline'):
            s = child.content
            if k == 'html_inline': s = re.sub(r'<[^>]+>', '', s)
            if not s: continue
            if link and re.match(r'https?://', link): external_link(par, s, link)
            else:
                r = par.add_run(s); r.bold = bold; r.italic = italic
                if k == 'code_inline': r.font.name = 'Consolas'
        elif k == 'image':
            uri = child.attrGet('src') or ''
            if re.match(r'https?://', uri):
                raise ValueError(f'External image not downloaded/embedded: {uri}. Store a real licensed asset locally and record provenance in MEDIA_MANIFEST.json')
            alternatives = [manuscript.parent / uri, manuscript.parent.parent / uri]
            pic = next((f for f in alternatives if f.is_file() and f.suffix.lower() in ('.jpg', '.jpeg', '.png', '.webp')), None)
            if pic is None: raise FileNotFoundError(f'Missing embedded photo/map file: {uri}')
            from PIL import Image
            with Image.open(pic) as im:
                w,h=im.size
                if w < 600 or h < 400: raise ValueError(f'Illustration too small to print: {pic.name} ({w}x{h})')
                if im.format not in ('PNG','JPEG'): raise ValueError(f'Convert image to PNG/JPEG for Word: {pic.name}')
            # Keep photos useful beside venue prose; full maps need more vertical room.
            max_w,max_h = 15.8, (12.5 if "maps" in pic.parts else 9.0)
            par.paragraph_format.keep_together = True
            par.paragraph_format.keep_with_next = True
            scale = min(max_w / w, max_h / h)
            par.add_run().add_picture(str(pic), width=Cm(w*scale), height=Cm(h*scale))
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif k == 'softbreak': par.add_run(' ')
        elif k == 'hardbreak': par.add_run('\n')


def cover_text(request, first_title):
    country = str(request.get('country') or '').strip() or 'TRAVEL'
    place = str(request.get('destination') or '').strip() or first_title
    return country.upper(), place.upper(), display_trip_dates(request)


def write_docx(markdown: Path, dest: Path, *, request=None, page_map=None):
    from markdown_it import MarkdownIt
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    content = markdown.read_text(encoding='utf-8')
    if len(content.strip()) < 1000: raise ValueError('Manuscript is not long enough for publication')
    tokens = MarkdownIt('commonmark').enable('table').parse(content)
    headings, cover_index = parse_headings(tokens)
    toc = [(idx, lev, name) for idx, (_, lev, name) in enumerate(headings) if lev <= 3]
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    sec.header_distance = Cm(0.85); sec.footer_distance = Cm(1.05)
    style_font(doc.styles['Normal'], size=10.5)
    # English / Latin defaults must also apply to bullet, table, footer and TOC text.
    style_font(doc.styles['Default Paragraph Font'], size=10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.2
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    for lev, sz in ((1, 19), (2, 14), (3, 11.5), (4, 10.5)):
        st = doc.styles[f'Heading {lev}']; style_font(st, size=sz, color=RGBColor(25, 25, 25))
        st.font.bold = True; st.paragraph_format.space_before = Pt(15 if lev < 3 else 8)
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.keep_with_next = True
    for lev in (1, 2, 3):
        nm = f'TOC {lev}'
        st = doc.styles[nm] if nm in doc.styles else doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        style_font(st, size=9.1 if lev < 3 else 8.8)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.left_indent = Cm((lev-1)*0.53)
        st.paragraph_format.keep_together = True
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(BRAND); r.bold = False; r.font.name = WESTERN_FONT; r.font.size = Pt(9)
    hpr = h._p.get_or_add_pPr(); borders = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '5'); bottom.set(qn('w:color'), '777777')
    borders.append(bottom); hpr.append(borders)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(f, ' PAGE ', '1'); f.add_run(' / '); add_field(f, ' NUMPAGES ', '1')
    settings = doc.settings.element
    auto = OxmlElement('w:updateFields'); auto.set(qn('w:val'), 'true'); settings.append(auto)

    # Cover (legacy editorial layout, not copied destination content)
    cover_from_md = tokens[cover_index+1].content if cover_index is not None else BRAND
    ctry, city, dates = cover_text(request or {}, cover_from_md)
    top = doc.add_paragraph(); top.paragraph_format.space_before = Pt(145)
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dot = top.add_run('●  '); dot.font.color.rgb = RGBColor(210, 9, 25); dot.font.size = Pt(23)
    r = top.add_run(ctry); r.font.size = Pt(33); r.font.bold = True; r.font.name = WESTERN_FONT
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    rr = p.add_run(city); rr.font.size = Pt(33); rr.font.bold = True; rr.font.name = WESTERN_FONT
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    rr = p.add_run(dates); rr.font.size = Pt(14); rr.font.name = WESTERN_FONT
    doc.add_page_break()

    # Three-level TOC with cached real page numbers, Word updatable TOC field,
    # dot leaders and internal bookmarks. Static cache works in headless PDF.
    t = doc.add_paragraph()
    label = t.add_run('Content'); label.font.name = WESTERN_FONT
    label.font.color.rgb = RGBColor(25, 25, 25); label.font.bold = True; label.font.size = Pt(21)
    t.paragraph_format.space_after = Pt(11)
    toc_field_open(doc.add_paragraph())
    for i, lev, name in toc:
        p = doc.add_paragraph(style=f'TOC {lev}')
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.4-(lev-1)*0.53), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        internal_link(p, name, i)
        p.add_run('\t' + str((page_map or {}).get(i, '000')))
    toc_field_close(doc.add_paragraph())

    # Markdown body with heading bookmarks, illustrations and linked sources.
    current = None; bullets = []; row = None; cell = None; table = None; col = 0
    heading_ids = {pos: idx for idx, (pos, _, _) in enumerate(headings)}
    skip = {cover_index, cover_index+1, cover_index+2} if cover_index is not None else set()
    for pos, tok in enumerate(tokens):
        if pos in skip: continue
        typ = tok.type
        if typ == 'heading_open':
            level = min(4, int(tok.tag[1])); current = doc.add_heading(level=level)
            if pos in heading_ids:
                bookmark(current, heading_ids[pos])
                if level in (1, 2): current.paragraph_format.page_break_before = True
        elif typ == 'paragraph_open':
            if cell is not None:
                current = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
            elif bullets:
                current = doc.add_paragraph(style='List Bullet' if bullets[-1] else 'List Number')
            else: current = doc.add_paragraph()
        elif typ == 'inline':
            if current is None: current = doc.add_paragraph()
            inline(current, tok, markdown)
        elif typ in ('bullet_list_open','ordered_list_open'): bullets.append(typ == 'bullet_list_open')
        elif typ in ('bullet_list_close','ordered_list_close'):
            if bullets: bullets.pop()
        elif typ in ('fence','code_block'):
            p = doc.add_paragraph(); r = p.add_run(tok.content.rstrip()); r.font.name='Consolas';r.font.size=Pt(8)
        elif typ == 'table_open':
            count=0
            for fut in tokens[pos+1:]:
                if fut.type=='tr_close': break
                if fut.type in ('td_open','th_open'): count += 1
            table=doc.add_table(rows=0, cols=max(1,count)); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
        elif typ=='tr_open': row=table.add_row();col=0
        elif typ in ('td_open','th_open'):
            cell=row.cells[col]; col+=1;current=cell.paragraphs[0]
        elif typ in ('td_close','th_close'): current=None;cell=None
        elif typ=='table_close':
            # Repeat headings and prohibit split rows when photos increase page flow.
            for idx, tr in enumerate(table.rows):
                pr = tr._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
                if idx == 0: pr.append(OxmlElement('w:tblHeader'))
            table=None;row=None;cell=None
        elif typ in ('heading_close','paragraph_close','blockquote_close'): current=None
    # An image paragraph keeps the immediately adjacent caption on the same page.
    previous_image = False
    for par in doc.paragraphs:
        has_image = bool(par._p.xpath('.//w:drawing'))
        if previous_image and not has_image and not par.style.name.startswith('Heading'):
            par.paragraph_format.keep_together = True
        previous_image = has_image
    # Table/hyperlink built-in styles can silently override the Normal font.
    for part in [doc.element, sec.header._element, sec.footer._element]:
        for run in part.xpath('.//w:r'):
            pr = run.find(qn('w:rPr'))
            if pr is None:
                pr = OxmlElement('w:rPr'); run.insert(0, pr)
            fonts = pr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts'); pr.insert(0, fonts)
            for key in ('ascii', 'hAnsi', 'cs'): fonts.set(qn('w:' + key), WESTERN_FONT)
            fonts.set(qn('w:eastAsia'), EAST_ASIAN_FONT)
    dest.parent.mkdir(parents=True, exist_ok=True); doc.save(dest)
    return toc


def pdf_from_docx(docx: Path, dest: Path):
    executable = shutil.which('libreoffice') or shutil.which('soffice')
    if executable is None: raise RuntimeError('LibreOffice / soffice unavailable; PDF and verified page-numbered TOC cannot be produced')
    with tempfile.TemporaryDirectory(prefix='periplus-lo-') as tmp:
        out = Path(tmp)
        profile = out / 'lo_profile'
        cmd = [executable, f'-env:UserInstallation=file://{profile}', '--headless',
               '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(out), str(docx)]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        candidate = out / (docx.stem + '.pdf')
        if result.returncode or not candidate.exists():
            raise RuntimeError('DOCX->PDF conversion failed: ' + (result.stderr or result.stdout)[:800])
        dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(candidate, dest)
    import fitz
    pdf = fitz.open(str(dest))
    if pdf.page_count < 3: raise RuntimeError('PDF lacks cover, contents and document body')
    absent = []
    for i, page in enumerate(pdf):
        header = ' '.join(b[4].replace('\n',' ') for b in page.get_text('blocks') if b[1] < 95)
        if BRAND not in header: absent.append(i+1)
    if absent: raise RuntimeError(f'Running header missing from PDF page(s): {absent[:10]}')
    # Cross-format typography gate. A DOCX can declare Oswald but the PDF may
    # silently substitute Noto/Liberation when Oswald is missing on this machine.
    font_names = {str(entry[3]).lower() for page in pdf for entry in page.get_fonts()}
    if not any('oswald' in name for name in font_names):
        raise RuntimeError(
            'PDF font verification failed: Oswald is not used. Install Oswald on '
            'the PDF conversion machine (fonts.google.com/specimen/Oswald), '
            'refresh the font cache / restart LibreOffice, then export again. '
            'Do not silently substitute another western font.'
        )
    return pdf


def outline_pages(pdf, toc):
    outline = pdf.get_toc()
    answers = {}; nxt = 0
    for idx, lev, name in toc:
        seek = name.strip()
        for j in range(nxt, len(outline)):
            a,b,page = outline[j][:3]
            if a == lev and b.strip() == seek:
                answers[idx] = page; nxt = j+1; break
        if idx not in answers: raise RuntimeError(f'PDF outline missing heading for TOC pagination: {name!r}')
    return answers


def audit_toc(pdf, toc, actual):
    first_section = min(actual.values()) if actual else None
    if first_section is None or first_section < 3:
        raise RuntimeError('Invalid body start page; cannot verify Contents')
    toc_text = '\n'.join(pdf[i].get_text() for i in range(1, first_section - 1))
    if 'Content' not in toc_text: raise RuntimeError('No Content page in PDF')
    # Cached entries must show actual page numbers, not zeros/placeholders.
    if '000' in toc_text: raise RuntimeError('TOC contains unfilled page number')
    for idx, lev, name in toc:
        if name not in toc_text:
            raise RuntimeError(f'TOC entry not rendered: {name}')
    last_page_text = pdf[-1].get_text()
    if not re.search(r'\d+\s*/\s*\d+', last_page_text):
        raise RuntimeError('Missing PAGE / NUMPAGES footer')


def export_project(project: Path, formats=None):
    formats = formats or ['docx','pdf']
    project = Path(project); source = project / 'exports' / f'{ROOT_NAME}.md'
    if not source.is_file(): raise FileNotFoundError(f'Missing researched manuscript: {source}')
    req_file = project/'request.json'
    request = json.loads(req_file.read_text(encoding='utf-8')) if req_file.is_file() else {}
    from reader_qa import audit_reader_text, audit_transport_plan
    problems = audit_reader_text(source.read_text(encoding='utf8'), request) + audit_transport_plan(project, source.read_text(encoding='utf8'), request)
    if problems: raise RuntimeError('Reader/transport QA failed: ' + '; '.join(problems[:12]))
    stem = publication_stem(request)
    docx = project/'exports'/f'{stem}.docx'
    pdf_file = project/'exports'/f'{stem}.pdf'
    try:
        import fitz
    except ImportError as ex:
        raise RuntimeError('Install pymupdf for deterministic numbered TOC and PDF QA') from ex
    pages = None
    for attempt in range(4):
        toc = write_docx(source, docx, request=request, page_map=pages)
        pdf = pdf_from_docx(docx, pdf_file)
        observed = outline_pages(pdf, toc)
        if pages == observed:
            audit_toc(pdf, toc, observed)
            pdf.close()
            break
        pdf.close(); pages = observed
    else:
        raise RuntimeError('TOC pagination did not stabilize after 4 iterations')
    # Publication invariant: real images and maps must be physically embedded, not just linked.
    from media_utils import verify_published_media
    verify_published_media(project, source, docx, pdf_file)
    results = []
    if 'docx' in formats: results.append(docx)
    if 'pdf' in formats: results.append(pdf_file)
    if 'md' in formats:
        public_md = project/'exports'/f'{stem}.md'
        shutil.copy2(source, public_md); results.append(public_md)
    return results


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--project',required=True)
    p.add_argument('--formats', nargs='+',choices=['docx','pdf','md'],default=['docx','pdf'])
    args=p.parse_args()
    for f in export_project(Path(args.project), args.formats): print(f)

if __name__=='__main__': main()
